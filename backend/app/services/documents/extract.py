"""Safe text extraction from untrusted uploads.

Defences
* Type is decided by CONTENT (magic bytes / container structure), never by the filename.
* Allowed: PDF, DOCX, UTF-8 plain text. Everything else is rejected (incl. legacy .doc,
  macro-enabled documents, images).
* DOCX: zip-bomb checks (entry count, total uncompressed size, compression ratio, path names)
  BEFORE parsing.
* Parsing runs in a separate process with a hard timeout and page/character caps, so a
  malformed or malicious file cannot hang or exhaust the API process.
* OCR is not performed: scanned PDFs without a text layer are reported as such.
"""

from __future__ import annotations

import io
import multiprocessing as mp
import zipfile
from dataclasses import dataclass

MAX_PDF_PAGES = 30
MAX_TEXT_CHARS = 120_000
DOCX_MAX_ENTRIES = 1000
DOCX_MAX_UNCOMPRESSED = 30 * 1024 * 1024
DOCX_MAX_RATIO = 120
EXTRACT_TIMEOUT_S = 20


class RejectedFile(ValueError):
    pass


@dataclass
class Extracted:
    mime: str
    ext: str
    text: str
    pages: int | None
    warnings: list[str]


def sniff(data: bytes) -> tuple[str, str]:
    if data.startswith(b"%PDF-"):
        return "application/pdf", "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile as exc:
            raise RejectedFile("Corrupt archive.") from exc
        if "[Content_Types].xml" in names and "word/document.xml" in names:
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise RejectedFile("Macro-enabled documents are not accepted.")
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
        raise RejectedFile("Only .docx Word documents are accepted among zip-based formats.")
    if b"\x00" in data[:4096]:
        raise RejectedFile("Unsupported binary file type. Upload PDF, DOCX or plain text.")
    try:
        data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise RejectedFile("Text files must be UTF-8 encoded.") from exc
    return "text/plain", "txt"


def _check_docx(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        if len(infos) > DOCX_MAX_ENTRIES:
            raise RejectedFile("Document has too many internal parts.")
        total = sum(i.file_size for i in infos)
        if total > DOCX_MAX_UNCOMPRESSED:
            raise RejectedFile("Document expands to an unsafe size.")
        for i in infos:
            if i.compress_size and i.file_size / max(i.compress_size, 1) > DOCX_MAX_RATIO:
                raise RejectedFile("Document has a suspicious compression ratio.")
            if i.filename.startswith(("/", "\\")) or ".." in i.filename.replace("\\", "/").split("/"):
                raise RejectedFile("Document contains unsafe internal paths.")


def _parse_worker(kind: str, data: bytes) -> tuple[str, int | None, list[str]]:
    warnings: list[str] = []
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise RejectedFile("Encrypted PDFs are not accepted.")
        n = len(reader.pages)
        if n > MAX_PDF_PAGES:
            warnings.append(f"Only the first {MAX_PDF_PAGES} of {n} pages were read.")
        parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            parts.append(page.extract_text() or "")
            if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                break
        text = "\n".join(parts)
        if len(text.strip()) < 40:
            warnings.append(
                "This PDF has little or no extractable text (it may be a scan). OCR is not enabled "
                "in this deployment; upload a text-based PDF or DOCX instead."
            )
        return text, n, warnings
    if kind == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines), None, warnings
    return data.decode("utf-8"), None, warnings


def _child(conn: object, kind: str, data: bytes) -> None:  # pragma: no cover - runs in subprocess
    try:
        conn.send(("ok", _parse_worker(kind, data)))  # type: ignore[attr-defined]
    except RejectedFile as exc:
        conn.send(("rejected", str(exc)))  # type: ignore[attr-defined]
    except Exception as exc:
        conn.send(("error", type(exc).__name__))  # type: ignore[attr-defined]


def extract(data: bytes, *, timeout: float = EXTRACT_TIMEOUT_S) -> Extracted:
    if not data:
        raise RejectedFile("The file is empty.")
    mime, ext = sniff(data)
    if ext == "docx":
        _check_docx(data)
    if ext == "txt":
        text, pages, warnings = _parse_worker("txt", data)
    else:
        ctx = mp.get_context("spawn")
        parent, child = ctx.Pipe(duplex=False)
        proc = ctx.Process(target=_child, args=(child, ext, data), daemon=True)
        proc.start()
        child.close()
        if not parent.poll(timeout):
            proc.kill()
            proc.join(5)
            raise RejectedFile("The document took too long to process and was rejected.")
        status, payload = parent.recv()
        proc.join(5)
        if status == "rejected":
            raise RejectedFile(str(payload))
        if status == "error":
            raise RejectedFile("The document could not be parsed (it may be corrupt or malformed).")
        text, pages, warnings = payload
    text = text.replace("\x00", "")[:MAX_TEXT_CHARS]
    return Extracted(mime=mime, ext=ext, text=text, pages=pages, warnings=warnings)
